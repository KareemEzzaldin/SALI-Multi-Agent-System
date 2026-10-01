"""
PPTX Parser for Lecture Presentations
Extracts slide titles, body text, speaker notes, tables, and isolates images.
"""
from __future__ import annotations
import io
import uuid
from typing import List, Tuple
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from ..models.ingestion_schemas import ExtractedSegment, ModalityType


class PPTXParser:
    @staticmethod
    def parse_pptx_bytes(pptx_bytes: bytes, file_name: str) -> Tuple[List[ExtractedSegment], List[dict]]:
        """
        Parses PPTX presentation bytes into structured segments and extracts slide images.
        """
        prs = Presentation(io.BytesIO(pptx_bytes))
        segments: List[ExtractedSegment] = []
        extracted_images: List[dict] = []

        for slide_index, slide in enumerate(prs.slides):
            slide_num = slide_index + 1
            slide_texts: List[str] = []
            title = ""

            for shape in slide.shapes:
                # 1. Slide Title
                if shape.has_text_frame and shape == slide.shapes.title:
                    title = shape.text.strip()

                # 2. Body Text Frames
                elif shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        p_text = paragraph.text.strip()
                        if p_text:
                            slide_texts.append(p_text)

                # 3. Tables
                elif shape.has_table:
                    table_rows: List[str] = []
                    for row in shape.table.rows:
                        row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                        table_rows.append("| " + " | ".join(row_cells) + " |")

                    if table_rows:
                        header_sep = "| " + " | ".join(["---"] * len(shape.table.columns)) + " |"
                        formatted_table = table_rows[0] + "\n" + header_sep + "\n" + "\n".join(table_rows[1:])
                        slide_texts.append("\n" + formatted_table + "\n")

                # 4. Embedded Pictures for Gemini OCR
                elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    try:
                        img = shape.image
                        img_bytes = img.blob
                        if len(img_bytes) > 5120:  # Skip tiny icons
                            extracted_images.append({
                                "page_num": slide_num,
                                "image_index": len(extracted_images) + 1,
                                "image_bytes": img_bytes,
                                "mime_type": img.content_type or "image/png",
                                "source_file": file_name
                            })
                    except Exception:
                        pass

            # 5. Speaker Notes (Crucial grounding context)
            speaker_notes = ""
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                speaker_notes = slide.notes_slide.notes_text_frame.text.strip()

            # Compile unified slide text
            full_content_parts = []
            if title:
                full_content_parts.append(f"## Slide {slide_num}: {title}")
            else:
                full_content_parts.append(f"## Slide {slide_num}")

            if slide_texts:
                full_content_parts.append("\n".join(slide_texts))

            if speaker_notes:
                full_content_parts.append(f"\n> **Instructor Speaker Notes:** {speaker_notes}")

            combined_text = "\n\n".join(full_content_parts).strip()

            if combined_text:
                segments.append(
                    ExtractedSegment(
                        segment_id=f"seg_{uuid.uuid4().hex[:8]}",
                        modality=ModalityType.SLIDE,
                        content=combined_text,
                        source_file=file_name,
                        page_or_slide_num=slide_num,
                        metadata={
                            "has_title": bool(title),
                            "has_notes": bool(speaker_notes),
                            "total_slides": len(prs.slides)
                        }
                    )
                )

        return segments, extracted_images
